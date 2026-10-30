#!/usr/bin/env python3
"""Build the faculty report from public evidence and an ignored private team file.

Requires Python with python-docx and Pillow installed.
Render and visually inspect the report before distribution.
"""
from pathlib import Path
import argparse
import json
import re
import shutil
import textwrap
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

root=Path(__file__).resolve().parents[1]
p=argparse.ArgumentParser();p.add_argument('--team',type=Path,required=True);p.add_argument('--output',type=Path,default=root/'output/submission');a=p.parse_args()
team=json.loads(a.team.read_text());out=a.output.resolve();out.mkdir(parents=True,exist_ok=True)
assets=root/'docs/submission/assets';assets.mkdir(parents=True,exist_ok=True)
schema=json.loads((root/'docs/submission/schema.json').read_text())
tables={t['name']:t for t in schema['tables']}
regular='/System/Library/Fonts/Supplemental/Arial.ttf';bold='/System/Library/Fonts/Supplemental/Arial Bold.ttf'
def font(n,b=False):return ImageFont.truetype(bold if b else regular,n)
ink='#172e2b';teal='#17695e';pale='#eef5f1';line='#9dbab0';gray='#52625d'

def wrap(draw,txt,width,f):
    words=txt.split();lines=[];s=''
    for word in words:
        if draw.textlength((s+' '+word).strip(),font=f)>width and s:lines.append(s);s=word
        else:s=(s+' '+word).strip()
    if s:lines.append(s)
    return lines

# Architecture is a deterministic code-native diagram.
im=Image.new('RGB',(2300,850),'white');d=ImageDraw.Draw(im)
boxes=[(50,300,450,520,'Browser','Researcher / Admin'),(570,300,1000,520,'Next.js','Vercel Hobby'),(1110,300,1530,520,'FastAPI','SQLAlchemy ORM'),(1690,70,2240,260,'PostgreSQL','System of record'),(1690,330,2240,520,'Verified model','CPU RoBERTa + GNN'),(1690,590,2240,780,'Neo4j','Graph projection')]
for x,y,x2,y2,title,sub in boxes:
 d.rounded_rectangle((x,y,x2,y2),radius=24,fill=pale,outline=line,width=3);d.text((x+26,y+38),title,font=font(46,True),fill=ink);d.text((x+26,y+110),sub,font=font(32),fill=gray)
for start,end in [((450,410),(570,410)),((1000,410),(1110,410)),((1530,410),(1690,165)),((1530,410),(1690,425)),((1530,410),(1690,685))]:
 d.line((start,end),fill=teal,width=6);ex,ey=end;d.polygon([(ex,ey),(ex-20,ey-12),(ex-20,ey+12)],fill=teal)
d.text((560,235),'HTTPS / same-origin proxy',font=font(30),fill=gray)
im.save(assets/'architecture.png',dpi=(300,300))

# Retrospective schedule uses visible repository/deployment milestone ranges.
im=Image.new('RGB',(2300,1200),'white');d=ImageDraw.Draw(im)
labels=['Repository and module setup','Data and model integration','Verified schema and evaluation','Cloud deployment and recovery','Workspace and final submission','Matched model evaluation']
phases=[(0,1),(2,17),(18,19),(22,22),(23,23),(25,25)]
left=720;top=180;cell=60;rh=135
for j,day in enumerate(range(6,32)):
 x=left+j*cell;d.text((x+10,top-85),str(day),font=font(25,True),fill=ink);d.line((x,top,x,top+len(labels)*rh),fill='#e4eae7',width=2)
d.text((left,35),'OCTOBER 2026',font=font(36,True),fill=teal)
for i,(label,(start,end)) in enumerate(zip(labels,phases)):
 y=top+i*rh;d.text((20,y+26),label,font=font(34),fill=ink);d.rounded_rectangle((left+start*cell+4,y+20,left+(end+1)*cell-4,y+85),radius=12,fill=teal)
d.text((20,1100),'Retrospective phases anchored in Git history and deployment records.',font=font(28),fill=gray)
im.save(assets/'gantt.png',dpi=(300,300))

# Domain ER plates contain every current table, with repeated context entities.
groups={
 'identity':['roles','users','session_tokens','reviewers','products','reviews','review_variants','review_feedback','audit_logs'],
 'model':['preprocessing_versions','model_versions','graph_snapshots','graph_snapshot_reviews','inference_runs','fraud_predictions','inductive_prediction_evidence','reviews'],
 'execution':['users','graph_snapshots','inference_runs','batch_analyses','batch_analysis_rows','inductive_inference_jobs','graph_sync_jobs','reviews'],
 'synthetic':['products','reviewers','reviews','synthetic_campaigns','synthetic_campaign_members','synthetic_campaign_reviews'],
}
covered=set()
for name,names in groups.items():
 covered.update(names);rows=(len(names)+1)//2;w=2300;h=100+rows*300
 im=Image.new('RGB',(w,h),'white');d=ImageDraw.Draw(im);positions={}
 for i,n in enumerate(names):
  col=i%2;row=i//2;positions[n]=(40+col*1160,60+row*300)
 # Light connectors behind cards; full field/cardinality definitions live in the dictionary.
 for n in names:
  x,y=positions[n]
  for c in tables[n]['columns']:
   for fk in c['foreign_keys']:
    parent=fk.split('.')[0]
    if parent in positions and parent!=n:
     px,py=positions[parent]
     start=(x+510,y+245);end=(px+510,py)
     d.line((start,end),fill='#b4c9c0',width=5)
     ex,ey=end;d.polygon([(ex,ey),(ex-12,ey+23),(ex+12,ey+23)],fill=teal)
 for n in names:
  x,y=positions[n];t=tables[n]
  d.rounded_rectangle((x,y,x+1020,y+245),radius=14,fill='white',outline=line,width=3)
  d.rounded_rectangle((x,y,x+1020,y+58),radius=14,fill=teal)
  d.text((x+20,y+11),n,font=font(33,True),fill='white')
  entries=[]
  for c in t['columns']:
   if c['pk']:entries.append('PK  '+c['name'])
  for c in t['columns']:
   if c['foreign_keys']:entries.append(('PK/FK  ' if c['pk'] else 'FK  ')+c['name']+' -> '+c['foreign_keys'][0].split('.')[0])
  if len(entries)<4:
   entries.extend(c['name'] for c in t['columns'] if c['name'] in ('status','source_system','version_key','role_key','snapshot_key'))
  if len(entries)>5:
   entries=entries[:4]+['Additional FK fields in full ERD and dictionary']
  yy=y+72
  for entry in entries:
   d.text((x+20,yy),entry,font=font(27),fill=ink);yy+=32
 d.text((40,h-26),'Child FK -> parent | Full attributes and nullable relationships: ERD.md and appendix',font=font(20),fill=gray)
 im.save(assets/f'erd-{name}.png',dpi=(300,300))
assert covered==set(tables),('Missing ERD entities',set(tables)-covered)
# Formal report, with editable text and a full dictionary appendix.
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.8);sec.left_margin=sec.right_margin=Inches(1)
for style in ('Normal','Body Text'):
 s=doc.styles[style];s.font.name='Arial';s.font.size=Pt(11);s.paragraph_format.space_after=Pt(7);s.paragraph_format.line_spacing=1.12
for style,size in [('Title',27),('Subtitle',14),('Heading 1',17),('Heading 2',13),('Heading 3',11)]:
 s=doc.styles[style];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
 s.paragraph_format.space_before=Pt(14);s.paragraph_format.space_after=Pt(9)
 s.paragraph_format.keep_with_next=True
for style in ('Heading 1','Heading 2','Heading 3'):doc.styles[style].font.bold=True
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('NeuroGraph-DB  |  ');r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def para(txt,style=None):
 p=doc.add_paragraph(txt,style);return p

def tab(headers,rows,widths=None,small=False):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths=widths or ([2.4,4.1] if headers==['Database view','Definition and use'] else [1.85,4.65] if len(headers)==2 else [6.5/len(headers)]*len(headers))
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 for i,h in enumerate(headers):t.rows[0].cells[i].text=h
 trpr=t.rows[0]._tr.get_or_add_trPr();rep=OxmlElement('w:tblHeader');trpr.append(rep)
 for row in rows:
  cells=t.add_row().cells
  for i,v in enumerate(row):cells[i].text=str(v)
 for ri,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'55' if small else '80');el.set(qn('w:type'),'dxa');m.append(el)
   pr.append(m)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   pr.append(borders)
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E5EBE8' if ri==0 else ('F6F8F7' if ri%2==0 else 'FFFFFF'));pr.append(sh)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.02
    for r in p.runs:r.font.size=Pt(9 if small else 10);r.bold=ri==0
 para('')
 return t

para('Vellore Institute of Technology').alignment=WD_ALIGN_PARAGRAPH.CENTER
para('School of Computer Science and Engineering').alignment=WD_ALIGN_PARAGRAPH.CENTER
para('Fall Semester 2026 to 2027').alignment=WD_ALIGN_PARAGRAPH.CENTER
para('')
p=para('NeuroGraph Database Systems Project','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=para('Final Project Submission','Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=para('A relational review investigation system with verified graph model inference');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
para('')
tab(['Team member','Registration number'],[[m[0],m[1]] for m in team['members']],[3.8,2.7])
tab(['Course details','Confirmed submission information'],[['Course',team['course']+'  '+team['code']],['Class number',team['class_number']],['Lab slot',team['slot']],['Course faculty',team['faculty']],['Submission date',datetime.now(ZoneInfo('Asia/Kolkata')).strftime('%d %B %Y')]])
para('Live application: https://frontend-three-snowy-10.vercel.app')
para('Private submission source: https://github.com/Avaneeshchavan/neurograph-db-submission')
doc.add_page_break()
para('Report guide','Heading 1')
for text in ['1  Project overview','2  Problem statement and objectives','3  Project planning and responsibilities','4  Technology and system architecture','5  Database design and ER diagrams','6  Normalization and index rationale','7  Database implementation','8  Application workflows and user experience','9  Authentication authorization and security','10 Backup recovery and operational evidence','11 Testing and deployment','12 Version control and bonus evidence','13 Faculty demonstration and submission contents','14 Limitations and future development','References','Appendix A Complete data dictionary','Appendix B Rubric coverage and demonstration evidence']:
 para(text)
para('The source and evidence files accompany this report. Access credentials and native database dumps are kept outside public source control.')
doc.add_page_break()

lines=(root/'docs/submission/REPORT.md').read_text().splitlines();i=0
while i<len(lines):
 s=lines[i].strip()
 if not s:i+=1;continue
 if s.startswith('# '):i+=1;continue
 if s.startswith('## '):
  title=s[3:]
  # Major sections begin on a clean page, avoiding squeezed headings after figures.
  if not title.startswith('1 Project') and title!='References':doc.add_page_break()
  para(title,'Heading 1');i+=1;continue
 if s.startswith('### '):para(s[4:],'Heading 2');i+=1;continue
 if s.startswith('!['):
  m=re.match(r'!\[(.*?)\]\((.*?)\)',s);path=root/'docs/submission'/m.group(2)
  is_erd='/erd-' in str(path)
  if 'live.png' in str(path) or 'live.jpg' in str(path):doc.add_page_break()
  if is_erd and not str(path).endswith('erd-identity.png'):doc.add_page_break()
  im=Image.open(path);maxw=6.5;maxh=7.5 if is_erd else 6.8
  width=min(maxw,maxh*im.width/im.height)
  pic=doc.add_paragraph();pic.alignment=WD_ALIGN_PARAGRAPH.CENTER;pic.add_run().add_picture(str(path),width=Inches(width))
  cap=para(m.group(1));cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
  for r in cap.runs:r.font.size=Pt(10);r.italic=True
  i+=1;continue
 if s.startswith('|'):
  data=[]
  while i<len(lines) and lines[i].strip().startswith('|'):
   row=[x.strip() for x in lines[i].strip().strip('|').split('|')]
   if not all(re.fullmatch(r':?-+:?',x) for x in row):data.append(row)
   i+=1
  if data:tab(data[0],data[1:])
  continue
 para(s);i+=1

# Complete physical field dictionary; no table is silently omitted.
doc.add_page_break();para('Appendix A Complete data dictionary','Heading 1')
para('Generated from SQLAlchemy metadata for migration 0007. PK denotes primary key; FK denotes referenced table and column. Composite uniqueness and CHECK definitions are also supplied in schema-postgresql.sql. The dictionary includes all 23 application tables; views and migration bookkeeping are not counted as domain tables.')
for name,t in tables.items():
 para(name.replace('_',' '),'Heading 2')
 rows=[]
 for c in t['columns']:
  rules=[]
  if c['pk']:rules.append('PK')
  rules+=['FK '+v for v in c['foreign_keys']]
  if c['default'] is not None:rules.append('Default '+c['default'])
  rows.append([c['name'],c['type'].replace('TIMESTAMP WITH TIME ZONE','TIMESTAMPTZ'),'Yes' if c['nullable'] else 'No','; '.join(rules) or '-'])
 tab(['Column','Type','Null','Key and default'],rows,[2.15,1.2,.45,2.7],small=True)
 checks=[c['expression'] for c in t['constraints'] if c['kind']=='CHECK']
 uniques=[', '.join(c.get('columns',[])) for c in t['constraints'] if c['kind']=='UNIQUE']
 if uniques:para('Unique constraints: '+'; '.join(uniques)+'.')
 if checks:
  para('CHECK constraints')
  for check in checks:
   pp=para(check)
   for rr in pp.runs:rr.font.size=Pt(9)
 if t['indexes']:
  para('Indexes: '+'; '.join(idx['name']+' ('+', '.join(idx['columns'])+')' for idx in t['indexes'])+'.')

doc.add_page_break();para('Appendix B Rubric coverage and demonstration evidence','Heading 1')
para('The rubric allocates 30 base marks and up to 3 bonus marks. Evidence is mapped below; this report does not guarantee faculty marks.')
rows=[['Planning','2','Problem, objectives, scope, responsibilities, stack, deliverables and dated Gantt'],['Database design','6','23-table ERD, relational schema, dictionary, keys, constraints and normalization analysis; physical provenance redundancy disclosed'],['Database implementation','8','ORM CRUD, joins, aggregate views, transaction rollback, stored procedure, trigger and indexes'],['Application','4','Framework dashboard, upload validation, saved results, search/filter/pagination, feedback CRUD and exceptions'],['Authentication and security','5','Single users table, roles FK, Argon2, JWT, refresh/logout, API RBAC and scoped data'],['Professional practices','5','Progressive Git/main/dev, private mirror, native restore, Docker and HTTPS cloud; Docker Hub publication not claimed'],['CI bonus','+1','GitHub Actions tests the real application and builds a Docker image'],['Custom domain bonus','Not claimed','Only provider hostnames are configured'],['Additional practice bonus','+1','Automated testing, redacted trigger audits, verified recovery and actual query-plan evidence']]
tab(['Area','Marks','Evidence and qualification'],rows,[1.5,.65,4.35])
para('Eight minute demonstration route','Heading 2')
for n,txt in enumerate(['Model results, metric scope and researcher CSV upload.','Saved history, reopen, identical retry and feedback CRUD.','Admin resolution, API denial, ERD and PostgreSQL features.','Native restore, Git main/dev, private collaboration and CI evidence.'],1):para(str(n)+'. '+txt)
para('Credentials are provided separately. Do not project passwords, private keys or native backup contents during the demonstration.')

doc.core_properties.title='NeuroGraph Database Systems Project'
doc.core_properties.subject='Database Systems Lab final project submission'
doc.core_properties.author='Avaneesh Chavan; Eldo John Giby; Anubhav Bhattacharya'
doc.core_properties.created=datetime.now(timezone.utc)
doc.core_properties.modified=doc.core_properties.created
doc.core_properties.comments=''
# Remove inherited title rules; formal report titles use whitespace only.
for element in list(doc.styles.element.iter(qn('w:pBdr'))) + list(doc.element.iter(qn('w:pBdr'))):
 element.getparent().remove(element)
path=out/'NeuroGraph_DB_Final_Report.docx';doc.save(path)
print('Created',path)
print('Complete dictionary tables',len(tables))
